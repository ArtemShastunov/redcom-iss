from io import BytesIO
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session, joinedload

from ..deps import get_current_user, get_db, require_role
from ..exporters.iss_xlsx import build_iss_xlsx_bytes
from ..models import Evaluation, Period, User, UserRole

router = APIRouter(prefix="/api/periods", tags=["export"])

MONTHS_RU = {
    1: "Январь",
    2: "Февраль",
    3: "Март",
    4: "Апрель",
    5: "Май",
    6: "Июнь",
    7: "Июль",
    8: "Август",
    9: "Сентябрь",
    10: "Октябрь",
    11: "Ноябрь",
    12: "Декабрь",
}

XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@router.get("/{period_id}/export")
def export_period_xlsx(
    period_id: int,
    db: Session = Depends(get_db),
    _: User = Depends(require_role(UserRole.SUPERVISOR, UserRole.HEAD, UserRole.ADMIN)),
) -> Response:
    """ПОЛНЫЙ экспорт за период (оценки всех операторов, статистика).
    Доступен только супервайзерам, начальнику ИСС и админу.
    """
    period = db.scalar(select(Period).where(Period.id == period_id))
    if period is None:
        raise HTTPException(404, "Period not found")

    month_label = MONTHS_RU.get(period.month, str(period.month))
    data = build_iss_xlsx_bytes(db, period, month_label)

    filename = f"Совместные итоги_{month_label}_{period.year}.xlsx"
    return Response(
        content=data,
        media_type=XLSX_MEDIA,
        headers={
            "Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}",
        },
    )


@router.get("/{period_id}/my-export")
def export_my_evaluations(
    period_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
) -> Response:
    """ЛИЧНЫЙ экспорт для оператора. Выгружает СТРОГО его оценки
    за период. Если у пользователя нет operator_id (например, он админ
    или начальник) — вернёт 400.
    """
    period = db.get(Period, period_id)
    if period is None:
        raise HTTPException(404, "Period not found")

    if user.operator_id is None:
        raise HTTPException(
            400,
            "У вашего аккаунта нет привязанного оператора. Обратитесь к администратору.",
        )

    evals = list(
        db.scalars(
            select(Evaluation)
            .options(joinedload(Evaluation.supervisor))
            .where(
                Evaluation.period_id == period_id,
                Evaluation.operator_id == user.operator_id,
            )
            .order_by(Evaluation.call_date, Evaluation.id)
        )
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Мои оценки"

    headers = [
        "Дата",
        "Телефон",
        "Договор",
        "Вопрос",
        "Тон",
        "Полнота",
        "1С",
        "Что было",
        "Как надо было",
        "На что повлияло",
        "Комментарий",
        "Супервайзер",
    ]
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")

    for row_idx, e in enumerate(evals, start=2):
        values = [
            e.call_date.date() if e.call_date else None,
            e.phone,
            e.contract_number,
            e.question,
            e.tone,
            e.completeness,
            e.dialog_1c,
            e.what_was,
            e.how_should_be,
            e.impact,
            e.comment,
            e.supervisor.full_name if e.supervisor else None,
        ]
        for col, v in enumerate(values, start=1):
            ws.cell(row=row_idx, column=col, value=v)

    widths = [12, 16, 14, 30, 6, 10, 6, 30, 30, 20, 30, 20]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    buf = BytesIO()
    wb.save(buf)

    month_label = MONTHS_RU.get(period.month, str(period.month))
    filename = f"Мои оценки_{user.full_name}_{month_label}_{period.year}.xlsx"

    return Response(
        content=buf.getvalue(),
        media_type=XLSX_MEDIA,
        headers={
            "Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}",
        },
    )
