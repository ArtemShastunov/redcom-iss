import tempfile
from io import BytesIO
from pathlib import Path
from urllib.parse import quote

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session, joinedload

from ..deps import get_current_user, get_db, require_role
from ..importers.debtor_xlsx import import_debtor_xlsx
from ..models import DebtorPeriod, DebtorRecord, User, UserRole
from ..schemas import (
    DebtorImportResult,
    DebtorPeriodOut,
    DebtorReportOut,
    DebtorRowOut,
)

router = APIRouter(prefix="/api/debtor", tags=["debtor"])

MONTHS_RU = {
    1: "Январь", 2: "Февраль", 3: "Март", 4: "Апрель",
    5: "Май", 6: "Июнь", 7: "Июль", 8: "Август",
    9: "Сентябрь", 10: "Октябрь", 11: "Ноябрь", 12: "Декабрь",
}

XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _compute_row(rec: DebtorRecord) -> DebtorRowOut:
    prev = float(rec.previous_amount or 0)
    curr = float(rec.current_amount or 0)
    total = prev + curr
    p5 = total * 0.05
    net = p5 * 0.87 * 1.6
    return DebtorRowOut(
        name=rec.name,
        operator_id=rec.operator_id,
        operator_name=rec.operator.full_name if rec.operator else None,
        is_operator=rec.operator_id is not None,
        previous_amount=round(prev, 2),
        current_amount=round(curr, 2),
        total=round(total, 2),
        percent_5=round(p5, 2),
        net_amount=round(net, 2),
    )


def _build_report(db: Session, period: DebtorPeriod) -> DebtorReportOut:
    records = list(db.scalars(
        select(DebtorRecord)
        .options(joinedload(DebtorRecord.operator))
        .where(DebtorRecord.period_id == period.id)
        .order_by(DebtorRecord.name)
    ))
    rows = [_compute_row(r) for r in records]

    return DebtorReportOut(
        period_id=period.id,
        year=period.year,
        month=period.month,
        rows=rows,
        total_previous=round(sum(r.previous_amount for r in rows), 2),
        total_current=round(sum(r.current_amount for r in rows), 2),
        total_total=round(sum(r.total for r in rows), 2),
        total_percent_5=round(sum(r.percent_5 for r in rows), 2),
        total_net=round(sum(r.net_amount for r in rows), 2),
    )


@router.get("/periods", response_model=list[DebtorPeriodOut])
def list_periods(
    db: Session = Depends(get_db),
    _: User = Depends(get_current_user),
) -> list[DebtorPeriod]:
    return list(db.scalars(
        select(DebtorPeriod).order_by(DebtorPeriod.year.desc(), DebtorPeriod.month.desc())
    ))


@router.get("/periods/{period_id}", response_model=DebtorReportOut)
def get_period(
    period_id: int,
    db: Session = Depends(get_db),
    _: User = Depends(get_current_user),
) -> DebtorReportOut:
    period = db.get(DebtorPeriod, period_id)
    if period is None:
        raise HTTPException(404, "Period not found")
    return _build_report(db, period)


@router.delete("/periods/{period_id}", status_code=204)
def delete_period(
    period_id: int,
    db: Session = Depends(get_db),
    _: User = Depends(require_role(UserRole.HEAD, UserRole.ADMIN)),
) -> None:
    period = db.get(DebtorPeriod, period_id)
    if period is None:
        raise HTTPException(404, "Period not found")
    db.delete(period)
    db.commit()


@router.post("/import", response_model=DebtorImportResult)
async def import_debtor(
    previous: UploadFile = File(...),
    current: UploadFile = File(...),
    year: int = 0,
    month: int = 0,
    db: Session = Depends(get_db),
    _: User = Depends(require_role(UserRole.HEAD, UserRole.ADMIN)),
) -> DebtorImportResult:
    if year < 2024 or year > 2100 or month < 1 or month > 12:
        raise HTTPException(400, "Укажите корректный год и месяц")

    for f in (previous, current):
        if not f.filename or not f.filename.lower().endswith(".xlsx"):
            raise HTTPException(400, "Оба файла должны быть .xlsx")

    with tempfile.TemporaryDirectory() as tmp:
        prev_path = Path(tmp) / "previous.xlsx"
        curr_path = Path(tmp) / "current.xlsx"
        prev_path.write_bytes(await previous.read())
        curr_path.write_bytes(await current.read())

        try:
            report = import_debtor_xlsx(
                previous_path=prev_path,
                current_path=curr_path,
                db=db,
                year=year,
                month=month,
            )
        except Exception as e:
            raise HTTPException(400, f"Ошибка импорта: {e}") from e

    return DebtorImportResult(
        period_id=report.period_id,
        year=year,
        month=month,
        matched=report.matched,
        unmatched=report.unmatched,
        previous_only=[],
        current_only=[],
    )


@router.get("/periods/{period_id}/export")
def export_period(
    period_id: int,
    db: Session = Depends(get_db),
    _: User = Depends(require_role(UserRole.HEAD, UserRole.ADMIN)),
) -> Response:
    period = db.get(DebtorPeriod, period_id)
    if period is None:
        raise HTTPException(404, "Period not found")

    report = _build_report(db, period)

    wb = Workbook()
    ws = wb.active
    ws.title = "Премия"

    headers = [
        "Оператор", "Остатки", MONTHS_RU.get(period.month, str(period.month)),
        "Итого", "5%", "На руки",
    ]
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")

    for row_idx, r in enumerate(report.rows, start=2):
        ws.cell(row=row_idx, column=1, value=r.name)
        ws.cell(row=row_idx, column=2, value=r.previous_amount)
        ws.cell(row=row_idx, column=3, value=r.current_amount)
        ws.cell(row=row_idx, column=4, value=r.total)
        ws.cell(row=row_idx, column=5, value=r.percent_5)
        ws.cell(row=row_idx, column=6, value=r.net_amount)

    total_row = len(report.rows) + 2
    ws.cell(row=total_row, column=1, value="ИТОГО").font = Font(bold=True)
    ws.cell(row=total_row, column=2, value=report.total_previous).font = Font(bold=True)
    ws.cell(row=total_row, column=3, value=report.total_current).font = Font(bold=True)
    ws.cell(row=total_row, column=4, value=report.total_total).font = Font(bold=True)
    ws.cell(row=total_row, column=5, value=report.total_percent_5).font = Font(bold=True)
    ws.cell(row=total_row, column=6, value=report.total_net).font = Font(bold=True)

    for i, w in enumerate([28, 16, 16, 16, 16, 16], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    buf = BytesIO()
    wb.save(buf)

    month_label = MONTHS_RU.get(period.month, str(period.month))
    filename = f"Премия_дебиторка_{month_label}_{period.year}.xlsx"

    return Response(
        content=buf.getvalue(),
        media_type=XLSX_MEDIA,
        headers={
            "Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}",
        },
    )
