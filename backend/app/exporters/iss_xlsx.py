"""Экспорт ведомости оценок звонков ИСС в Excel.

Формат повторяет «Совместные итоги_<Месяц>.xlsx»:

  * лист на каждого оператора: A1 = заголовок колонок, A2 = ФИО оператора,
    B2 = первый супервайзер, далее 15 звонков первого супервайзера,
    строка с именем второго супервайзера, 15 звонков второго,
    строка «Итого» с Excel-формулами (INT(AVERAGE));
  * лист «Статистика» — сводка с премиями;
  * лист «Детализация» — авторасчёт, override, место, премия.
"""
from __future__ import annotations

from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..models import (
    Evaluation, Operator, OperatorPeriod, Period,
    SupervisorPeriod, User, UserRole,
)
from ..scoring import (
    award_supervisor,
    rank_and_award_operators,
    score_operator,
)

HEADERS = [
    "ФИО", "", "Номер телефона", "Дата звонка",
    "Тон разговора", "Полнота ответа", "Диалог в 1С",
    "Номер договора", "Вопрос клиента", "Что было",
    "Как надо было", "На что повлияло", "Дополнительные комментарии",
]

HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
BOLD = Font(bold=True)


# ---------- примитивы ----------

def _write_row(ws, row: int, values: list) -> None:
    for col, v in enumerate(values, start=1):
        ws.cell(row=row, column=col, value=v)


def _style_header_row(ws, row: int, n_cols: int) -> None:
    for col in range(1, n_cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = BOLD
        cell.fill = HEADER_FILL


def _eval_to_row(e: Evaluation) -> list:
    return [
        None, None,
        e.phone, e.call_date,
        e.tone, e.completeness, e.dialog_1c,
        e.contract_number,
        e.question, e.what_was, e.how_should_be, e.impact, e.comment,
    ]


def _write_operator_sheet(
    ws,
    operator: Operator,
    evals: list[Evaluation],
    supervisors_order: list[str],
) -> None:
    _write_row(ws, 1, HEADERS)
    _style_header_row(ws, 1, len(HEADERS))

    ws.cell(row=2, column=1, value=operator.full_name)
    if supervisors_order:
        ws.cell(row=2, column=2, value=supervisors_order[0])

    row = 3
    for idx, sup_name in enumerate(supervisors_order):
        if idx > 0:
            ws.cell(row=row, column=2, value=sup_name)
            row += 1

        chunk = [e for e in evals if e.supervisor.full_name == sup_name]
        for e in chunk:
            _write_row(ws, row, _eval_to_row(e))
            row += 1

    ws.cell(row=row, column=1, value="Итого:").font = BOLD
    start = 3
    end = row - 1
    ws.cell(row=row, column=5, value=f"=INT(AVERAGE(E{start}:E{end}))")
    ws.cell(row=row, column=6, value=f"=INT(AVERAGE(F{start}:F{end}))")
    ws.cell(row=row, column=7, value=f"=INT(AVERAGE(G{start}:G{end}))")
    ws.cell(row=row, column=8, value=f"=INT((E{row}+F{row}+G{row})/3)")

    widths = [20, 8, 16, 18, 8, 10, 10, 14, 30, 40, 40, 20, 40]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _write_statistics_sheet(
    ws,
    month_label: str,
    operators_summary: list,
    supervisors_rows: list[dict],
    total_premium: int,
) -> None:
    _write_row(ws, 1, [month_label, "Тон", "", "", "Полнота ответа", "", "",
                       "1С и письма", "", "Общий итог", ""])
    _write_row(ws, 2, ["", "-", "+", "Оц.", "-", "+", "Оц.",
                       "-", "Оц.", "Оц.", "Сумма премии"])
    _style_header_row(ws, 2, 11)

    row = 3
    for op in operators_summary:
        _write_row(ws, row, [
            op.operator_name,
            None, None, op.summary_score,
            None, None, op.summary_score,
            None, op.summary_score,
            op.summary_score, op.premium,
        ])
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="Супервайзеры").font = BOLD
    row += 1
    for s in supervisors_rows:
        _write_row(ws, row, [
            s["name"],
            None, None, s["norm"],
            None, None, None, None, None, None,
            s["premium"],
        ])
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="Всего").font = BOLD
    ws.cell(row=row, column=11, value=total_premium).font = BOLD

    widths = [25, 6, 6, 8, 14, 6, 6, 14, 6, 8, 14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _write_details_sheet(ws, operators_summary: list) -> None:
    headers = [
        "Оператор", "Звонков", "Оценок",
        "Авто: тон", "Авто: полнота", "Авто: диалог",
        "Итог: тон", "Итог: полнота", "Итог: диалог",
        "Сводная", "Override: причина",
        "Место", "Премия", "Исключён",
    ]
    _write_row(ws, 1, headers)
    _style_header_row(ws, 1, len(headers))

    row = 2
    for op in operators_summary:
        _write_row(ws, row, [
            op.operator_name,
            op.calls_handled,
            op.evaluations_count,
            op.auto_tone, op.auto_completeness, op.auto_dialog,
            op.tone_score, op.completeness_score, op.dialog_score,
            op.summary_score,
            op.override_reason,
            op.rank,
            op.premium,
            op.excluded_reason,
        ])
        row += 1

    widths = [22, 8, 8, 10, 12, 12, 10, 12, 12, 8, 30, 6, 8, 20]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ---------- публичное API ----------

def build_iss_xlsx_bytes(db: Session, period: Period, month_label: str) -> bytes:
    """Возвращает xlsx-файл в памяти. Ничего не пишет на диск."""
    evals = list(db.scalars(
        select(Evaluation)
        .where(Evaluation.period_id == period.id)
        .order_by(Evaluation.operator_id, Evaluation.call_date, Evaluation.id)
    ))
    by_op: dict[int, list[Evaluation]] = {}
    for e in evals:
        by_op.setdefault(e.operator_id, []).append(e)

    operators = list(db.scalars(select(Operator).order_by(Operator.id)))
    operators_by_id = {op.id: op for op in operators}

    op_periods = {
        op.operator_id: op
        for op in db.scalars(
            select(OperatorPeriod).where(OperatorPeriod.period_id == period.id)
        )
    }

    supervisors = list(db.scalars(
        select(User).where(User.role == UserRole.SUPERVISOR).order_by(User.id)
    ))
    supervisors_order = [s.full_name for s in supervisors]

    summaries = []
    for op_id, op in operators_by_id.items():
        evs = by_op.get(op_id, [])
        op_period = op_periods.get(op_id)
        if not evs and (op_period is None or op_period.calls_handled == 0):
            continue
        calls = op_period.calls_handled if op_period else 0
        if calls == 0:
            calls = len(evs)
        summaries.append(score_operator(
            operator_id=op.id, operator_name=op.full_name,
            tone_scores=[e.tone for e in evs],
            completeness_scores=[e.completeness for e in evs],
            dialog_scores=[e.dialog_1c for e in evs],
            calls_handled=calls,
            has_critical_errors=op_period.has_critical_errors if op_period else False,
            critical_errors_reason=op_period.critical_errors_reason if op_period else None,
            op_period=op_period,
        ))
    summaries = rank_and_award_operators(summaries)

    sup_periods = {
        sp.supervisor_id: sp
        for sp in db.scalars(
            select(SupervisorPeriod).where(SupervisorPeriod.period_id == period.id)
        )
    }
    sup_rows: list[dict] = []
    for s in supervisors:
        sp = sup_periods.get(s.id)
        norm = sp.norm_score if sp else None
        sup_rows.append({
            "name": s.full_name,
            "norm": norm,
            "premium": award_supervisor(norm),
        })

    total_premium = sum(o.premium for o in summaries) + sum(r["premium"] for r in sup_rows)

    wb = Workbook()

    first = True
    for op_id, op in operators_by_id.items():
        evs = by_op.get(op_id, [])
        if not evs:
            continue
        if first:
            ws = wb.active
            ws.title = op.full_name[:31]
            first = False
        else:
            ws = wb.create_sheet(title=op.full_name[:31])
        _write_operator_sheet(ws, op, evs, supervisors_order)

    if first:
        ws = wb.active
        ws.title = "Пусто"
        ws.cell(row=1, column=1, value="Нет данных за период")

    ws_stat = wb.create_sheet(title="Статистика")
    _write_statistics_sheet(ws_stat, month_label, summaries, sup_rows, total_premium)

    ws_det = wb.create_sheet(title="Детализация")
    _write_details_sheet(ws_det, summaries)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_iss_xlsx(
    path: Path,
    db: Session,
    period: Period,
    month_label: str,
) -> None:
    data = build_iss_xlsx_bytes(db, period, month_label)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)
