import argparse
import getpass
import sys
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select, text

from .db import SessionLocal
from .exporters.iss_xlsx import export_iss_xlsx
from .importers.iss_xlsx import import_iss_xlsx
from .importers.callbacks_xlsx import import_callbacks_xlsx
from .importers.premiums_xlsx import import_premiums_xlsx
from .models import Operator, Period, User, UserRole
from .security import hash_password, make_login_from_full_name

MONTHS_RU = {
    1: "Январь",
    2: "Февраль",
    3: "Март",
    4: "Апрель",
    5: "Май",
    6: "Июнь",
    7: "Июль",
    8: "Август",
    9: "Сентябрь",
    10: "Октябрь",
    11: "Ноябрь",
    12: "Декабрь",
}


def _parse_fio(s: str) -> tuple[str, str | None]:
    parts = s.strip().split()
    if not parts:
        return "", None
    return parts[0], " ".join(parts[1:]) if len(parts) > 1 else None


# ---------- команды ----------


def cmd_create_user(args: argparse.Namespace) -> int:
    db = SessionLocal()
    try:
        if db.scalar(select(User).where(User.login == args.login)):
            print(f"User {args.login!r} already exists", file=sys.stderr)
            return 1
        password = args.password or getpass.getpass("Password: ")
        if not password:
            print("Empty password", file=sys.stderr)
            return 1
        role = UserRole(args.role)
        db.add(
            User(
                login=args.login,
                password_hash=hash_password(password),
                full_name=args.full_name,
                role=role,
            )
        )
        db.commit()
        print(f"Created {role.value} {args.login!r}")
        return 0
    finally:
        db.close()


def cmd_change_password(args: argparse.Namespace) -> int:
    db = SessionLocal()
    try:
        user = db.scalar(select(User).where(User.login == args.login))
        if user is None:
            print(f"User {args.login!r} not found", file=sys.stderr)
            return 1
        password = args.password or getpass.getpass("New password: ")
        if not password:
            print("Empty password", file=sys.stderr)
            return 1
        user.password_hash = hash_password(password)
        db.commit()
        print(f"Password changed for {args.login!r}")
        return 0
    finally:
        db.close()


def cmd_create_operator_accounts(args: argparse.Namespace) -> int:
    """Массово создаёт аккаунты операторам, у которых их ещё нет."""
    default_password = args.password or "changeme2026"

    db = SessionLocal()
    try:
        operators = list(
            db.scalars(
                select(Operator)
                .where(Operator.is_active.is_(True))
                .order_by(Operator.id)
            )
        )

        created: list[tuple[str, str, str]] = []
        skipped: list[tuple[str, str]] = []

        for op in operators:
            existing = db.scalar(
                select(User).where(
                    User.operator_id == op.id,
                    User.role == UserRole.OPERATOR,
                )
            )
            if existing is not None:
                skipped.append((op.full_name, existing.login))
                continue

            base_login = make_login_from_full_name(op.full_name)
            candidate = base_login
            n = 2
            while db.scalar(select(User).where(User.login == candidate)):
                candidate = f"{base_login}_{n}"
                n += 1

            user = User(
                login=candidate,
                password_hash=hash_password(default_password),
                full_name=op.full_name,
                role=UserRole.OPERATOR,
                is_active=True,
                operator_id=op.id,
            )
            db.add(user)
            created.append((op.full_name, candidate, default_password))

        db.commit()

        print(f"Создано аккаунтов: {len(created)}")
        if created:
            print()
            print(f"{'ФИО':<32}{'Логин':<20}Пароль")
            print("-" * 72)
            for full_name, login, pwd in created:
                print(f"{full_name:<32}{login:<20}{pwd}")

        if skipped:
            print()
            print(f"Пропущено (уже есть аккаунт): {len(skipped)}")
            for full_name, login in skipped:
                print(f"  {full_name}: {login}")

        return 0
    finally:
        db.close()


def cmd_import_iss_xlsx(args: argparse.Namespace) -> int:
    db = SessionLocal()
    try:
        report = import_iss_xlsx(
            path=Path(args.path),
            db=db,
            year=args.year,
            month=args.month,
            replace=not args.no_replace,
        )
    except Exception as e:
        print(f"Import failed: {e}", file=sys.stderr)
        return 2
    finally:
        db.close()

    print(f"Period: {args.year}-{args.month:02d} (id={report.period_id})")
    print(f"Evaluations deleted: {report.evaluations_deleted}")
    print(f"Evaluations created: {report.evaluations_created}")
    print(f"Operators created:   {len(report.operators_created)}")
    print(f"Operators reused:    {len(report.operators_reused)}")
    print(f"Supervisors created: {len(report.supervisors_created)}")
    print(f"Supervisors reused:  {len(report.supervisors_reused)}")
    if report.per_operator_counts:
        print("Per-operator evaluations:")
        for name, n in sorted(report.per_operator_counts.items()):
            print(f"  {name}: {n}")
    if report.warnings:
        print("Warnings:")
        for w in report.warnings:
            print(f"  - {w}")
    return 0


def cmd_export_iss_xlsx(args: argparse.Namespace) -> int:
    db = SessionLocal()
    try:
        period = db.scalar(
            select(Period).where(Period.year == args.year, Period.month == args.month)
        )
        if period is None:
            print(f"Period {args.year}-{args.month:02d} not found", file=sys.stderr)
            return 1
        month_label = MONTHS_RU.get(args.month, str(args.month))
        export_iss_xlsx(
            path=Path(args.path),
            db=db,
            period=period,
            month_label=month_label,
        )
    except Exception as e:
        print(f"Export failed: {e}", file=sys.stderr)
        return 2
    finally:
        db.close()

    print(f"Exported to {args.path}")
    return 0


def cmd_fill_premium_names(args: argparse.Namespace) -> int:
    path = Path(args.path)
    if not path.exists():
        print(f"File not found: {path}", file=sys.stderr)
        return 1
    if path.suffix.lower() != ".xlsx":
        print(f"Only .xlsx is supported (got {path.suffix!r})", file=sys.stderr)
        return 1

    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    entries: list[tuple[str, str]] = []
    for row_idx in range(2, ws.max_row + 1):
        raw = ws.cell(row=row_idx, column=2).value
        if raw is None:
            continue
        raw = str(raw).strip()
        if not raw or raw.lower().startswith("итого"):
            continue
        surname, _ = _parse_fio(raw)
        if surname:
            entries.append((surname, raw))

    by_surname: dict[str, list[str]] = {}
    for surname, raw in entries:
        by_surname.setdefault(surname.lower(), []).append(raw)

    db = SessionLocal()
    try:
        operators = list(db.scalars(select(Operator).order_by(Operator.id)))
        updated: list[tuple[str, str]] = []
        already_set: list[str] = []
        ambiguous: list[tuple[str, list[str]]] = []
        not_in_file: list[str] = []

        for op in operators:
            if op.premium_name and not args.force:
                already_set.append(op.full_name)
                continue
            op_surname, _ = _parse_fio(op.full_name)
            matches = by_surname.get(op_surname.lower(), [])
            if len(matches) == 1:
                op.premium_name = matches[0]
                updated.append((op.full_name, matches[0]))
            elif len(matches) == 0:
                not_in_file.append(op.full_name)
            else:
                ambiguous.append((op.full_name, matches))

        db.commit()
        print(f"Updated: {len(updated)}")
        for db_name, prem in updated:
            print(f"  {db_name}  ->  {prem}")
        if already_set:
            print(f"Already set: {len(already_set)}")
        if ambiguous:
            print(f"Ambiguous: {len(ambiguous)}")
        if not_in_file:
            print(f"Not found in file: {len(not_in_file)}")
        return 0
    finally:
        db.close()


def cmd_import_premiums_xlsx(args: argparse.Namespace) -> int:
    db = SessionLocal()
    try:
        report = import_premiums_xlsx(
            path=Path(args.path),
            db=db,
            year=args.year,
            month=args.month,
        )
    except Exception as e:
        print(f"Import failed: {e}", file=sys.stderr)
        return 2
    finally:
        db.close()

    print(f"Period: {args.year}-{args.month:02d} (id={report.period_id})")
    print(f"Rows in file: {report.rows_total}")
    print(f"Matched:      {len(report.matched)}")
    if report.created:
        print(f"OperatorPeriod created: {report.created}")
    if report.updated:
        print(f"OperatorPeriod updated: {report.updated}")
    if report.not_matched:
        print(f"Not matched: {len(report.not_matched)}")
        for raw in report.not_matched:
            print(f"  {raw}")
    if report.ambiguous:
        print(f"Ambiguous: {len(report.ambiguous)}")
    return 0


def cmd_import_callbacks_xlsx(args: argparse.Namespace) -> int:
    db = SessionLocal()
    try:
        report = import_callbacks_xlsx(
            path=Path(args.path),
            db=db,
            year=args.year,
            month=args.month,
        )
    except Exception as e:
        print(f"Import failed: {e}", file=sys.stderr)
        return 2
    finally:
        db.close()

    print(f"Period: {args.year}-{args.month:02d} (id={report.period_id})")
    print(f"Rows in file: {report.rows_total}")
    print(f"Matched:      {len(report.matched)}")
    for db_name, file_name, count in report.matched:
        print(f"  {db_name}  <-  {file_name}: {count}")
    if report.created:
        print(f"OperatorPeriod created: {report.created}")
    if report.updated:
        print(f"OperatorPeriod updated: {report.updated}")
    if report.not_matched:
        print(f"Not matched: {len(report.not_matched)}")
        for name, count in report.not_matched:
            print(f"  {name}: {count}")
    return 0


def cmd_delete_period(args: argparse.Namespace) -> int:
    db = SessionLocal()
    try:
        period = db.scalar(
            select(Period).where(Period.year == args.year, Period.month == args.month)
        )
        if period is None:
            print(f"Period not found", file=sys.stderr)
            return 1
        pid = period.id
        db.execute(text("DELETE FROM evaluations WHERE period_id = :pid"), {"pid": pid})
        db.execute(
            text("DELETE FROM operator_periods WHERE period_id = :pid"), {"pid": pid}
        )
        db.execute(
            text("DELETE FROM supervisor_periods WHERE period_id = :pid"), {"pid": pid}
        )
        db.execute(text("DELETE FROM periods WHERE id = :pid"), {"pid": pid})
        db.commit()
        print(f"Deleted period {args.year}-{args.month:02d} (id={pid})")
        return 0
    finally:
        db.close()


def cmd_reset_db(args: argparse.Namespace) -> int:
    db = SessionLocal()
    try:
        if not args.yes:
            answer = input("Delete all data? [y/N]: ").strip().lower()
            if answer != "y":
                print("Aborted")
                return 1
        db.execute(text("DELETE FROM evaluations"))
        db.execute(text("DELETE FROM operator_periods"))
        db.execute(text("DELETE FROM supervisor_periods"))
        db.execute(text("DELETE FROM periods"))
        db.execute(text("DELETE FROM operators"))
        if args.with_users:
            db.execute(text("DELETE FROM users"))
        db.commit()
        print("Done")
        return 0
    finally:
        db.close()


def main() -> int:
    parser = argparse.ArgumentParser(prog="app.cli")
    sub = parser.add_subparsers(dest="cmd", required=True)

    p = sub.add_parser("create-user")
    p.add_argument("--login", required=True)
    p.add_argument("--full-name", required=True)
    p.add_argument("--role", required=True, choices=[r.value for r in UserRole])
    p.add_argument("--password", default=None)
    p.set_defaults(func=cmd_create_user)

    p = sub.add_parser("change-password")
    p.add_argument("--login", required=True)
    p.add_argument("--password", default=None)
    p.set_defaults(func=cmd_change_password)

    p = sub.add_parser(
        "create-operator-accounts",
        help="Массово создать аккаунты всем активным операторам",
    )
    p.add_argument(
        "--password",
        default=None,
        help="Пароль для всех (по умолчанию changeme2026)",
    )
    p.set_defaults(func=cmd_create_operator_accounts)

    p = sub.add_parser("import-iss-xlsx")
    p.add_argument("path")
    p.add_argument("--year", type=int, required=True)
    p.add_argument("--month", type=int, required=True)
    p.add_argument("--no-replace", action="store_true")
    p.set_defaults(func=cmd_import_iss_xlsx)

    p = sub.add_parser("export-iss-xlsx")
    p.add_argument("path")
    p.add_argument("--year", type=int, required=True)
    p.add_argument("--month", type=int, required=True)
    p.set_defaults(func=cmd_export_iss_xlsx)

    p = sub.add_parser("fill-premium-names")
    p.add_argument("path")
    p.add_argument("--force", action="store_true")
    p.set_defaults(func=cmd_fill_premium_names)

    p = sub.add_parser("import-premiums-xlsx")
    p.add_argument("path")
    p.add_argument("--year", type=int, required=True)
    p.add_argument("--month", type=int, required=True)
    p.set_defaults(func=cmd_import_premiums_xlsx)

    p = sub.add_parser("import-callbacks-xlsx")
    p.add_argument("path")
    p.add_argument("--year", type=int, required=True)
    p.add_argument("--month", type=int, required=True)
    p.set_defaults(func=cmd_import_callbacks_xlsx)

    p = sub.add_parser("delete-period")
    p.add_argument("--year", type=int, required=True)
    p.add_argument("--month", type=int, required=True)
    p.set_defaults(func=cmd_delete_period)

    p = sub.add_parser("reset-db")
    p.add_argument("--with-users", action="store_true")
    p.add_argument("--yes", action="store_true")
    p.set_defaults(func=cmd_reset_db)

    args = parser.parse_args()
    return args.func(args)


if __name__ == "__main__":
    raise SystemExit(main())
