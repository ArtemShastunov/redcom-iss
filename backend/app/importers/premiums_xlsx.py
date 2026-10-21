from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..models import Operator, OperatorPeriod, Period, PeriodStatus


@dataclass
class PremiumRow:
    raw_name: str
    incoming: int
    outgoing: int
    bitrix: int
    mobile: int


@dataclass
class PremiumsImportReport:
    period_id: int
    rows_total: int = 0
    matched: list[tuple[str, str]] = field(default_factory=list)
    not_matched: list[str] = field(default_factory=list)
    ambiguous: list[tuple[str, list[str]]] = field(default_factory=list)
    created: int = 0
    updated: int = 0


def _parse_fio(s: str) -> tuple[str, str | None]:
    parts = s.strip().split()
    if not parts:
        return "", None
    return parts[0], " ".join(parts[1:]) if len(parts) > 1 else None


def _to_int(v) -> int:
    if v is None:
        return 0
    try:
        return int(v)
    except (TypeError, ValueError):
        return 0


def _parse_premium_file(path: Path) -> list[PremiumRow]:
    if path.suffix.lower() != ".xlsx":
        raise ValueError(f"Только .xlsx поддерживается (получено {path.suffix!r})")

    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    rows = []
    for row_idx in range(2, ws.max_row + 1):
        raw_name = ws.cell(row=row_idx, column=2).value
        if raw_name is None:
            continue
        raw_name = str(raw_name).strip()
        if not raw_name or raw_name.lower().startswith("итого"):
            continue
        rows.append(
            PremiumRow(
                raw_name=raw_name,
                incoming=_to_int(ws.cell(row=row_idx, column=3).value),
                outgoing=_to_int(ws.cell(row=row_idx, column=4).value),
                bitrix=_to_int(ws.cell(row=row_idx, column=5).value),
                mobile=_to_int(ws.cell(row=row_idx, column=6).value),
            )
        )
    return rows


def _build_indexes(
    operators: list[Operator],
) -> tuple[dict[str, Operator], dict[str, list[Operator]]]:
    by_premium: dict[str, Operator] = {}
    by_surname: dict[str, list[Operator]] = {}
    for op in operators:
        if op.premium_name:
            by_premium[op.premium_name.strip().lower()] = op
        surname, _ = _parse_fio(op.full_name)
        by_surname.setdefault(surname.lower(), []).append(op)
    return by_premium, by_surname


def _find_operator(
    raw_name: str,
    by_premium: dict[str, Operator],
    by_surname: dict[str, list[Operator]],
) -> tuple[Operator | None, str | None]:
    exact = by_premium.get(raw_name.strip().lower())
    if exact is not None:
        return exact, None
    surname, _ = _parse_fio(raw_name)
    matches = by_surname.get(surname.lower(), [])
    if len(matches) == 1:
        return matches[0], None
    if len(matches) > 1:
        return None, f"ambiguous: {[m.full_name for m in matches]}"
    return None, None


def import_premiums_xlsx(
    path: Path, db: Session, year: int, month: int
) -> PremiumsImportReport:
    period = db.scalar(select(Period).where(Period.year == year, Period.month == month))
    if period is None:
        period = Period(year=year, month=month, status=PeriodStatus.OPEN)
        db.add(period)
        db.flush()

    report = PremiumsImportReport(period_id=period.id)
    rows = _parse_premium_file(path)
    report.rows_total = len(rows)

    # ВАЖНО: Сбрасываем старые звонки и чаты за этот период в 0
    op_periods_all = list(
        db.scalars(select(OperatorPeriod).where(OperatorPeriod.period_id == period.id))
    )
    for op in op_periods_all:
        op.calls_incoming = 0
        op.calls_outgoing = 0
        op.chats_bitrix = 0
        op.chats_mobile = 0

    operators = list(db.scalars(select(Operator).order_by(Operator.id)))
    by_premium, by_surname = _build_indexes(operators)

    for row in rows:
        op, ambiguity = _find_operator(row.raw_name, by_premium, by_surname)
        if op is None:
            if ambiguity:
                report.ambiguous.append((row.raw_name, [ambiguity]))
            else:
                report.not_matched.append(row.raw_name)
            continue

        op_period = db.get(OperatorPeriod, (period.id, op.id))
        if op_period is None:
            op_period = OperatorPeriod(period_id=period.id, operator_id=op.id)
            db.add(op_period)
            report.created += 1
        else:
            report.updated += 1

        op_period.calls_incoming = row.incoming
        op_period.calls_outgoing = row.outgoing
        op_period.chats_bitrix = row.bitrix
        op_period.chats_mobile = row.mobile
        report.matched.append((op.full_name, row.raw_name))

    db.commit()
    return report
