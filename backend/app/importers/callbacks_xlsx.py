from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..models import Operator, OperatorPeriod, Period, PeriodStatus
from .premiums_xlsx import _build_indexes, _find_operator


@dataclass
class CallbacksImportReport:
    period_id: int
    rows_total: int = 0
    matched: list[tuple[str, str, int]] = field(default_factory=list)
    not_matched: list[tuple[str, int]] = field(default_factory=list)
    created: int = 0
    updated: int = 0


def _parse_callbacks_file(path: Path) -> dict[str, int]:
    if path.suffix.lower() != ".xlsx":
        raise ValueError(f"Только .xlsx поддерживается (получено {path.suffix!r})")

    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb.worksheets[0]

    counts: dict[str, int] = {}
    for row_idx in range(2, ws.max_row + 1):
        operator_raw = ws.cell(row=row_idx, column=6).value  # F - оператор
        if operator_raw is None:
            continue
        name = str(operator_raw).strip()
        if not name:
            continue
        counts[name] = counts.get(name, 0) + 1
    return counts


def import_callbacks_xlsx(
    path: Path, db: Session, year: int, month: int
) -> CallbacksImportReport:
    period = db.scalar(select(Period).where(Period.year == year, Period.month == month))
    if period is None:
        period = Period(year=year, month=month, status=PeriodStatus.OPEN)
        db.add(period)
        db.flush()

    report = CallbacksImportReport(period_id=period.id)
    counts = _parse_callbacks_file(path)
    report.rows_total = sum(counts.values())

    # ВАЖНО: Сбрасываем старые заказные звонки за этот период в 0
    op_periods_all = list(
        db.scalars(select(OperatorPeriod).where(OperatorPeriod.period_id == period.id))
    )
    for op in op_periods_all:
        op.callback_calls = 0

    operators = list(db.scalars(select(Operator).order_by(Operator.id)))
    by_premium, by_surname = _build_indexes(operators)

    for file_name, count in counts.items():
        op, _ = _find_operator(file_name, by_premium, by_surname)
        if op is None:
            report.not_matched.append((file_name, count))
            continue

        op_period = db.get(OperatorPeriod, (period.id, op.id))
        if op_period is None:
            op_period = OperatorPeriod(period_id=period.id, operator_id=op.id)
            db.add(op_period)
            report.created += 1
        else:
            report.updated += 1

        op_period.callback_calls = count
        report.matched.append((op.full_name, file_name, count))

    db.commit()
    return report
