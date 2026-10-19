"""Импорт ведомости оценок звонков ИСС из Excel.

Формат — «Совместные итоги_<Месяц>.xlsx»:
  * операторские листы (Layout A): A1 = "ФИО", A2 = ФИО оператора,
    в теле — блоки по 15 звонков, каждый блок начинается со строки,
    где в колонке A или B стоит имя супервайзера;
  * листы супервайзеров (Layout B): лист назван по имени супервайзера
    («София», «Ольга»), в теле — плоский список строк без привязки
    к оператору;
  * лист «Статистика» игнорируется.

Идемпотентен: повторный запуск полностью пересоздаёт оценки периода.

При импорте применяются псевдонимы имён (см. OPERATOR_NAME_ALIASES) —
так «Голубятников Леша» из старых файлов превращается в
«Голубятников Алексей», и не создаёт дубликат в БД.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import delete, select
from sqlalchemy.orm import Session

from ..models import (
    Evaluation, Operator, Period, PeriodStatus,
    User, UserRole,
)
from ..security import hash_password

SUPERVISOR_NAMES = {"София", "Ольга"}
SKIP_SHEETS = {"Статистика"}
HEADER_CELL = "ФИО"

STUB_PASSWORD = "changeme2026"
STUB_LOGIN_BY_NAME = {
    "София": "sofia",
    "Ольга": "olga",
}

# Псевдонимы имён. Ключ — как в файле Excel, значение — как в БД.
# Ключи храним в нижнем регистре и с обрезанными пробелами.
OPERATOR_NAME_ALIASES: dict[str, str] = {
    "голубятников леша": "Голубятников Алексей",
}


def _normalize_operator_name(name: str) -> str:
    key = name.strip().lower()
    return OPERATOR_NAME_ALIASES.get(key, name.strip())


@dataclass
class ImportReport:
    period_id: int
    operators_created: list[str] = field(default_factory=list)
    operators_reused: list[str] = field(default_factory=list)
    supervisors_created: list[str] = field(default_factory=list)
    supervisors_reused: list[str] = field(default_factory=list)
    evaluations_created: int = 0
    evaluations_deleted: int = 0
    warnings: list[str] = field(default_factory=list)
    per_operator_counts: dict[str, int] = field(default_factory=dict)


# ---------- утилиты ----------

def _cell_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, str):
        return v.strip()
    return str(v).strip()


def _cell_score(v) -> int:
    if v is None:
        raise ValueError("empty score cell")
    try:
        n = int(v)
    except (TypeError, ValueError):
        raise ValueError(f"invalid score value: {v!r}")
    if n not in (-1, 0, 1):
        raise ValueError(f"score out of range: {n!r}")
    return n


def _cell_date(v) -> datetime:
    if isinstance(v, datetime):
        return v
    if isinstance(v, str):
        for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d.%m.%Y"):
            try:
                return datetime.strptime(v.strip(), fmt)
            except ValueError:
                continue
    raise ValueError(f"invalid date: {v!r}")


def _is_phone_row(sheet, row_idx: int) -> bool:
    v = sheet.cell(row=row_idx, column=3).value
    if v is None:
        return False
    s = str(v).strip()
    if not s or s.startswith("="):
        return False
    if s.lower().startswith("итого"):
        return False
    return True


# ---------- сущности ----------

def _get_or_create_supervisor(db: Session, name: str, report: ImportReport) -> User:
    login = STUB_LOGIN_BY_NAME.get(name, name.lower())
    user = db.scalar(select(User).where(User.login == login))
    if user is not None:
        if name not in report.supervisors_reused:
            report.supervisors_reused.append(name)
        return user
    user = User(
        login=login,
        password_hash=hash_password(STUB_PASSWORD),
        full_name=name,
        role=UserRole.SUPERVISOR,
    )
    db.add(user)
    db.flush()
    report.supervisors_created.append(name)
    return user


def _get_or_create_operator(db: Session, full_name: str, report: ImportReport) -> Operator:
    normalized = _normalize_operator_name(full_name)
    op = db.scalar(select(Operator).where(Operator.full_name == normalized))
    if op is not None:
        if normalized not in report.operators_reused:
            report.operators_reused.append(normalized)
        return op
    op = Operator(full_name=normalized, is_active=True)
    db.add(op)
    db.flush()
    report.operators_created.append(normalized)
    return op


def _get_or_create_period(db: Session, year: int, month: int) -> Period:
    p = db.scalar(select(Period).where(Period.year == year, Period.month == month))
    if p is not None:
        return p
    p = Period(year=year, month=month, status=PeriodStatus.CLOSED,
               closed_at=datetime.utcnow())
    db.add(p)
    db.flush()
    return p


# ---------- парсинг листов ----------

def _read_layout_a(sheet, db: Session, period: Period, report: ImportReport) -> int:
    operator_name = _cell_str(sheet.cell(row=2, column=1).value)
    if not operator_name:
        raise ValueError(f"[{sheet.title}] пустая ячейка A2: не найден оператор")
    operator = _get_or_create_operator(db, operator_name, report)

    current_supervisor: User | None = None
    b2 = _cell_str(sheet.cell(row=2, column=2).value)
    if b2 in SUPERVISOR_NAMES:
        current_supervisor = _get_or_create_supervisor(db, b2, report)

    created = 0

    for row_idx in range(3, sheet.max_row + 1):
        a_val = _cell_str(sheet.cell(row=row_idx, column=1).value)
        b_val = _cell_str(sheet.cell(row=row_idx, column=2).value)

        candidate = None
        for c in (a_val, b_val):
            if c in SUPERVISOR_NAMES:
                candidate = c
                break

        if candidate is not None:
            current_supervisor = _get_or_create_supervisor(db, candidate, report)
            continue

        if a_val.lower().startswith("итого") or b_val.lower().startswith("итого"):
            break

        if not _is_phone_row(sheet, row_idx):
            continue

        if current_supervisor is None:
            report.warnings.append(
                f"[{sheet.title}] строка {row_idx}: оценка без активного супервайзера, пропущена"
            )
            continue

        ev = Evaluation(
            period_id=period.id,
            operator_id=operator.id,
            supervisor_id=current_supervisor.id,
            phone=_cell_str(sheet.cell(row=row_idx, column=3).value),
            call_date=_cell_date(sheet.cell(row=row_idx, column=4).value),
            tone=_cell_score(sheet.cell(row=row_idx, column=5).value),
            completeness=_cell_score(sheet.cell(row=row_idx, column=6).value),
            dialog_1c=_cell_score(sheet.cell(row=row_idx, column=7).value),
            contract_number=_cell_str(sheet.cell(row=row_idx, column=8).value) or None,
            question=_cell_str(sheet.cell(row=row_idx, column=9).value) or None,
            what_was=_cell_str(sheet.cell(row=row_idx, column=10).value) or None,
            how_should_be=_cell_str(sheet.cell(row=row_idx, column=11).value) or None,
            impact=_cell_str(sheet.cell(row=row_idx, column=12).value) or None,
            comment=_cell_str(sheet.cell(row=row_idx, column=13).value) or None,
        )
        db.add(ev)
        created += 1

    report.per_operator_counts[operator.full_name] = (
        report.per_operator_counts.get(operator.full_name, 0) + created
    )
    return created


def _read_layout_b(sheet, db: Session, period: Period, report: ImportReport) -> int:
    a1 = _cell_str(sheet.cell(row=1, column=1).value)
    a2 = _cell_str(sheet.cell(row=2, column=1).value)
    supervisor_name = (
        sheet.title if sheet.title in SUPERVISOR_NAMES
        else a1 if a1 in SUPERVISOR_NAMES
        else a2 if a2 in SUPERVISOR_NAMES
        else None
    )
    if supervisor_name is None:
        raise ValueError(
            f"[{sheet.title}] не удалось определить супервайзера "
            f"(A1={a1!r}, A2={a2!r})"
        )
    _get_or_create_supervisor(db, supervisor_name, report)
    report.warnings.append(
        f"[{sheet.title}] лист супервайзера пропущен: "
        f"в файле нет колонки с оператором для этих оценок"
    )
    return 0


# ---------- публичный API ----------

def import_iss_xlsx(
    path: Path,
    db: Session,
    year: int,
    month: int,
    replace: bool = True,
) -> ImportReport:
    if not path.exists():
        raise FileNotFoundError(path)

    report = ImportReport(period_id=0)
    period = _get_or_create_period(db, year, month)
    report.period_id = period.id

    if replace:
        deleted = db.execute(
            delete(Evaluation).where(Evaluation.period_id == period.id)
        ).rowcount or 0
        report.evaluations_deleted = deleted
        db.flush()

    wb = load_workbook(path, data_only=False, read_only=True)

    for sheet_name in wb.sheetnames:
        if sheet_name in SKIP_SHEETS:
            continue
        sheet = wb[sheet_name]
        a1 = _cell_str(sheet.cell(row=1, column=1).value)
        a2 = _cell_str(sheet.cell(row=2, column=1).value)

        if sheet_name in SUPERVISOR_NAMES:
            report.evaluations_created += _read_layout_b(sheet, db, period, report)
            continue

        if a1 == HEADER_CELL and a2:
            report.evaluations_created += _read_layout_a(sheet, db, period, report)
            continue

        if a1 in SUPERVISOR_NAMES or a2 in SUPERVISOR_NAMES:
            report.evaluations_created += _read_layout_b(sheet, db, period, report)
            continue

        report.warnings.append(
            f"[{sheet_name}] не опознан layout (A1={a1!r}, A2={a2!r}), пропущен"
        )

    db.commit()
    return report
